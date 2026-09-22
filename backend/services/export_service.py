"""
Export Service — generates Excel, PDF, and CSV files from exam schedules.
"""

from __future__ import annotations
import csv
import io
from sqlalchemy.orm import Session
from models.exam_schedule import ExamSchedule


def _build_rows(db: Session, schedules: list[ExamSchedule]) -> list[dict]:
    """Build flat dicts from schedule objects for export."""
    rows = []
    for s in schedules:
        rows.append({
            "Date": s.exam_date,
            "Time": f"{s.time_slot.start_time} - {s.time_slot.end_time}" if s.time_slot else "",
            "Class": s.section.parent_class.name if s.section and s.section.parent_class else "",
            "Section": s.section.name if s.section else "",
            "Batch": s.batch.name if s.batch else "",
            "Subject": s.subject.name if s.subject else "",
            "Venue": s.laboratory.name if s.laboratory else "",
            "Main In-Charge (Coordinator)": s.incharge.name if s.incharge else "",
            "Co-In-Charge": s.co_incharge.name if s.co_incharge else "",

            "Status": s.status,
        })
    return rows


def export_csv(db: Session, schedules: list[ExamSchedule]) -> str:
    """Return CSV string of the schedule."""
    rows = _build_rows(db, schedules)
    if not rows:
        return ""
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=rows[0].keys())
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


def export_excel_bytes(db: Session, schedules: list[ExamSchedule]) -> bytes:
    """Return Excel file bytes of the schedule."""
    from openpyxl import Workbook

    rows = _build_rows(db, schedules)
    wb = Workbook()
    ws = wb.active
    ws.title = "Exam Schedule"

    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for r in rows:
            ws.append([r[h] for h in headers])

        # Style header row
        from openpyxl.styles import Font, PatternFill
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="337AB7", end_color="337AB7", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        # Auto-width columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_pdf_bytes(db: Session, schedules: list[ExamSchedule]) -> bytes:
    """Return PDF file bytes of the schedule."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    rows = _build_rows(db, schedules)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    elements = []

    styles = getSampleStyleSheet()
    elements.append(Paragraph("PRACTICAL EXAMINATION SCHEDULE", styles["Title"]))
    elements.append(Spacer(1, 20))

    if rows:
        headers = list(rows[0].keys())
        table_data = [headers]
        for r in rows:
            table_data.append([str(r[h]) for h in headers])

        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#337AB7")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FC")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    doc.build(elements)
    return buffer.getvalue()
